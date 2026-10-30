"""Tableau de bord, chiffres (URSSAF, douanes) et tableurs téléchargeables."""
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Response, abort, flash, redirect, render_template, request, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from calculs import STATUTS, cout_revient, marge_pct, taux_tva, to_float
from noyau import (MODES_PAIEMENT, MOIS, TYPES_CLIENTS, app, couts_par_lot, get_db,
                   lots_avec_stock, mois_label, parse_date, settings, stock_lot, today)
from rappels import ca_periode, calculer_rappels, periodes_urssaf


# --------------------------------------------------------------------------
# Tableau de bord
# --------------------------------------------------------------------------

def stats_tickets(conn, conf, debut, fin, couts):
    tva = taux_tva(conf)
    row = conn.execute(
        """SELECT COUNT(*) n, COALESCE(SUM(sous_total - remise), 0) ca FROM tickets
           WHERE date_vente BETWEEN ? AND ?""", (debut, fin)).fetchone()
    lignes = conn.execute(
        "SELECT lot_id, quantite FROM ventes WHERE date_vente BETWEEN ? AND ?",
        (debut, fin)).fetchall()
    cout = sum(l["quantite"] * couts.get(l["lot_id"], 0) for l in lignes)
    ht = row["ca"] / (1 + tva)
    return {"ca": row["ca"], "tickets": row["n"], "btl": sum(l["quantite"] for l in lignes),
            "marge": marge_pct(ht, cout) if ht else None}


@app.route("/")
def tableau_de_bord():
    conn = get_db()
    conf = settings()
    couts = couts_par_lot(conn, conf)
    ajd = date.today()
    debut_mois = ajd.replace(day=1).isoformat()
    mois = stats_tickets(conn, conf, debut_mois, today(), couts)
    annee = stats_tickets(conn, conf, (ajd - timedelta(days=365)).isoformat(), today(), couts)

    lots = lots_avec_stock()
    lots_en_stock = [l for l in lots if stock_lot(l) > 0]

    historique = []
    y, m = ajd.year, ajd.month
    for _ in range(6):
        ym = f"{y}-{m:02d}"
        ca = conn.execute("SELECT COALESCE(SUM(sous_total - remise), 0) FROM tickets"
                          " WHERE date_vente LIKE ?", (ym + "%",)).fetchone()[0]
        historique.append({"label": MOIS[m - 1], "ca": ca})
        y, m = (y, m - 1) if m > 1 else (y - 1, 12)
    historique.reverse()

    prochains = conn.execute(
        """SELECT l.*, COALESCE(l.date_embouteillage,
                  date(l.date_maceration, '+' || COALESCE(r.duree_maceration_jours, 60) || ' days'))
                  AS date_prete
           FROM lots l LEFT JOIN recettes r ON r.id = l.recette_id
           WHERE l.statut IN ('maceration', 'pret')
           ORDER BY l.statut = 'pret' DESC, date_prete"""
    ).fetchall()
    evenements = conn.execute(
        "SELECT * FROM evenements WHERE date_debut >= ? ORDER BY date_debut LIMIT 3",
        (today(),)).fetchall()
    dernieres = conn.execute(
        """SELECT t.*, (SELECT SUM(quantite) FROM ventes v WHERE v.ticket_id = t.id) AS btl
           FROM tickets t ORDER BY t.date_vente DESC, t.id DESC LIMIT 5""").fetchall()

    return render_template(
        "tableau_de_bord.html", mois=mois, annee=annee, mois_label=mois_label(today()[:7]),
        stock_total=sum(stock_lot(l) for l in lots), lots_en_stock=lots_en_stock,
        historique=historique, ca_max=max(h["ca"] for h in historique) or 1,
        prochains=prochains, evenements=evenements, dernieres=dernieres,
        rappels=calculer_rappels(conn, conf), stock_lot=stock_lot,
    )


# --------------------------------------------------------------------------
# Chiffres : URSSAF et douanes
# --------------------------------------------------------------------------

@app.route("/chiffres")
def chiffres():
    conn = get_db()
    conf = settings()
    periodicite = conf.get("urssaf_periodicite", "mensuelle")
    taux = to_float(conf.get("urssaf_taux"))
    cfp = to_float(conf.get("cfp_taux"))
    periodes = []
    for cle, libelle, debut, fin, echeance in periodes_urssaf(
            periodicite, 12 if periodicite == "mensuelle" else 4):
        ca = ca_periode(conn, debut, fin)
        decl = conn.execute("SELECT * FROM declarations WHERE periode = ?",
                            (f"urssaf:{cle}",)).fetchone()
        nb = conn.execute("SELECT COUNT(*) FROM tickets WHERE date_vente BETWEEN ? AND ?",
                          (debut.isoformat(), fin.isoformat())).fetchone()[0]
        periodes.append({"cle": cle, "libelle": libelle, "debut": debut.isoformat(),
                         "fin": fin.isoformat(), "echeance": echeance.isoformat(), "ca": ca,
                         "cotisations": ca * taux / 100, "cfp": ca * cfp / 100,
                         "tickets": nb, "declaration": decl})
    choisie = next((p for p in periodes if p["cle"] == request.args.get("periode")),
                   periodes[0] if periodes else None)

    # Douanes : volumes d'alcool pur entrés (embouteillés) et sortis (vendus), par mois
    douanes = []
    y, m = date.today().year, date.today().month
    for _ in range(12):
        ym = f"{y}-{m:02d}"
        sorties = conn.execute(
            """SELECT COALESCE(SUM(v.quantite), 0) btl,
                      COALESCE(SUM(v.quantite * l.volume_cl / 100.0), 0) litres,
                      COALESCE(SUM(v.quantite * l.volume_cl / 100.0 * COALESCE(l.degre, 0) / 100), 0) lap
               FROM ventes v JOIN lots l ON l.id = v.lot_id WHERE v.date_vente LIKE ?""",
            (ym + "%",)).fetchone()
        entrees = conn.execute(
            """SELECT COALESCE(SUM(nb_bouteilles), 0) btl,
                      COALESCE(SUM(nb_bouteilles * volume_cl / 100.0 * COALESCE(degre, 0) / 100), 0) lap
               FROM lots WHERE statut IN ('embouteille', 'epuise') AND date_embouteillage LIKE ?""",
            (ym + "%",)).fetchone()
        douanes.append({"mois": mois_label(ym), "sorties": sorties, "entrees": entrees})
        y, m = (y, m - 1) if m > 1 else (y - 1, 12)
    stock_lap = sum(stock_lot(l) * l["volume_cl"] / 100 * (l["degre"] or 0) / 100
                    for l in lots_avec_stock())
    annee = date.today().year
    return render_template(
        "chiffres.html", periodes=periodes, choisie=choisie, periodicite=periodicite,
        taux=taux, cfp=cfp, douanes=douanes, stock_lap=stock_lap,
        debut_annee=f"{annee}-01-01", fin_annee=f"{annee}-12-31",
    )


@app.route("/chiffres/declarer", methods=["POST"])
def chiffres_declarer():
    conn = get_db()
    periode = request.form.get("periode", "")
    if not periode:
        abort(400)
    if request.form.get("action") == "annuler":
        conn.execute("DELETE FROM declarations WHERE periode = ?", (f"urssaf:{periode}",))
        flash("Déclaration marquée comme non faite.", "info")
    else:
        conn.execute("INSERT OR REPLACE INTO declarations (periode, declare_le, montant)"
                     " VALUES (?, ?, ?)",
                     (f"urssaf:{periode}", today(), to_float(request.form.get("montant"))))
        flash("Déclaration enregistrée comme faite. Le rappel disparaît du tableau de bord.",
              "succes")
    conn.commit()
    return redirect(url_for("chiffres", periode=periode))


# --------------------------------------------------------------------------
# Tableurs Excel
# --------------------------------------------------------------------------

ENTETE = PatternFill("solid", fgColor="0F1D33")
EURO = '#,##0.00 "€"'


def _feuille(wb, titre, colonnes, lignes, formats=None):
    """colonnes : [(titre, largeur)] ; formats : {index_colonne: format}."""
    ws = wb.create_sheet(titre[:31])
    ws.append([c for c, _ in colonnes])
    for i, (_, largeur) in enumerate(colonnes, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="E6C983")
        cell.fill = ENTETE
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largeur
    for ligne in lignes:
        ws.append([date.fromisoformat(v) if isinstance(v, str) and len(v) == 10
                   and v[4] == "-" and v[7] == "-" else v for v in ligne])
    for idx, fmt in (formats or {}).items():
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                cell.number_format = fmt
    for i, (titre_col, _) in enumerate(colonnes, 1):
        if titre_col.lower().startswith("date") or titre_col in ("Macération", "Embouteillage",
                                                                   "DLC", "Dernier achat",
                                                                   "Expiration", "Émission"):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for cell in row:
                    cell.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return ws


def feuille_livre(wb, conn, debut, fin):
    rows = conn.execute(
        """SELECT t.*, (SELECT group_concat(v.quantite || ' × ' || l.recette_nom, ' ; ')
                        FROM ventes v JOIN lots l ON l.id = v.lot_id WHERE v.ticket_id = t.id)
                       AS detail
           FROM tickets t WHERE t.date_vente BETWEEN ? AND ? ORDER BY t.date_vente, t.numero""",
        (debut, fin)).fetchall()
    _feuille(wb, "Livre des recettes",
             [("Date", 12), ("N° ticket", 14), ("N° facture", 14), ("Client", 28),
              ("Nature de la vente", 46), ("Mode de règlement", 16), ("Montant encaissé", 16),
              ("dont consigne", 13), ("CA (hors consigne)", 17)],
             [(t["date_vente"], t["numero"], t["facture_numero"] or "", t["client_nom"],
               t["detail"] or "", MODES_PAIEMENT.get(t["mode_paiement"], ""), t["total"],
               t["consigne"], t["sous_total"] - t["remise"]) for t in rows],
             {7: EURO, 8: EURO, 9: EURO})


def feuille_ventes(wb, conn, debut, fin, couts):
    rows = conn.execute(
        """SELECT v.*, t.numero AS ticket, t.mode_paiement, l.numero AS lot, l.recette_nom,
                  e.nom AS evenement
           FROM ventes v JOIN tickets t ON t.id = v.ticket_id JOIN lots l ON l.id = v.lot_id
           LEFT JOIN evenements e ON e.id = t.evenement_id
           WHERE v.date_vente BETWEEN ? AND ? ORDER BY v.date_vente, v.id""",
        (debut, fin)).fetchall()
    _feuille(wb, "Ventes détaillées",
             [("Date", 12), ("Ticket", 14), ("Client", 26), ("Canal", 11), ("Lot", 14),
              ("Produit", 30), ("Quantité", 10), ("Prix unitaire", 13), ("Total", 12),
              ("Coût de revient", 15), ("Marge", 12), ("Paiement", 12), ("Événement", 26)],
             [(v["date_vente"], v["ticket"], v["client"], v["canal"], v["lot"], v["recette_nom"],
               v["quantite"], v["prix_unitaire_ttc"], v["quantite"] * v["prix_unitaire_ttc"],
               v["quantite"] * couts.get(v["lot_id"], 0),
               v["quantite"] * (v["prix_unitaire_ttc"] - couts.get(v["lot_id"], 0)),
               MODES_PAIEMENT.get(v["mode_paiement"], ""), v["evenement"] or "") for v in rows],
             {8: EURO, 9: EURO, 10: EURO, 11: EURO})


def feuille_lots(wb, conn, conf):
    lots = lots_avec_stock("ORDER BY l.numero")
    _feuille(wb, "Registre des lots",
             [("Lot", 14), ("Produit", 30), ("Statut", 18), ("Macération", 12),
              ("Embouteillage", 13), ("DLC", 12), ("Bouteilles", 11), ("Contenance (cl)", 14),
              ("Degré", 8), ("Vendues", 9), ("En stock", 9), ("Coût / bouteille", 15),
              ("Prix vente TTC", 14), ("Notes", 40)],
             [(l["numero"], l["recette_nom"], STATUTS[l["statut"]], l["date_maceration"],
               l["date_embouteillage"], l["dlc"], l["nb_bouteilles"], l["volume_cl"], l["degre"],
               l["vendu"], stock_lot(l), cout_revient(conn, l, conf)["total"],
               l["prix_vente_ttc"], l["notes"] or "") for l in lots],
             {12: EURO, 13: EURO})
    ing = conn.execute(
        """SELECT l.numero, i.* FROM lot_ingredients i JOIN lots l ON l.id = i.lot_id
           ORDER BY l.numero, i.id""").fetchall()
    _feuille(wb, "Ingrédients des lots",
             [("Lot", 14), ("Catégorie", 12), ("Ingrédient", 32), ("Quantité", 10),
              ("Unité", 8), ("Coût", 11)],
             [(i["numero"], i["categorie"], i["nom"], i["quantite"], i["unite"], i["cout"])
              for i in ing], {6: EURO})


def feuille_stock(wb, conn, conf, couts):
    lots = [l for l in lots_avec_stock("ORDER BY l.numero") if stock_lot(l) > 0]
    _feuille(wb, "Stock bouteilles",
             [("Lot", 14), ("Produit", 30), ("En stock", 10), ("Coût unitaire", 13),
              ("Valeur", 12), ("DLC", 12)],
             [(l["numero"], l["recette_nom"], stock_lot(l), couts.get(l["id"], 0),
               stock_lot(l) * couts.get(l["id"], 0), l["dlc"]) for l in lots],
             {4: EURO, 5: EURO})
    mats = conn.execute("SELECT * FROM materiaux ORDER BY type, nom").fetchall()
    _feuille(wb, "Stock matières",
             [("Matière", 34), ("Type", 12), ("Stock", 9), ("Seuil d'alerte", 13),
              ("Coût unitaire", 13), ("À commander", 12)],
             [(m["nom"], m["type"], m["stock"], m["seuil_alerte"], m["cout_unitaire"],
               "OUI" if m["stock"] <= m["seuil_alerte"] else "") for m in mats], {5: EURO})


def feuille_matieres(wb, conn):
    rows = conn.execute(
        """SELECT m.*, COALESCE((SELECT SUM(quantite) FROM mouvements_mp v WHERE v.mp_id = m.id), 0)
           AS stock FROM matieres_premieres m ORDER BY m.categorie, m.nom""").fetchall()
    _feuille(wb, "Stock matières premières",
             [("Matière", 32), ("Catégorie", 11), ("Stock", 10), ("Unité", 8), ("Seuil", 8),
              ("Prix unitaire", 13), ("Valeur", 12), ("Fournisseur", 26), ("À racheter", 11)],
             [(m["nom"], m["categorie"], round(m["stock"], 3), m["unite"], m["seuil_alerte"],
               m["cout_unitaire"], max(m["stock"], 0) * m["cout_unitaire"], m["fournisseur"],
               "OUI" if m["seuil_alerte"] and m["stock"] <= m["seuil_alerte"] else "")
              for m in rows], {6: EURO, 7: EURO})
    mvts = conn.execute(
        """SELECT v.*, m.nom, m.unite FROM mouvements_mp v JOIN matieres_premieres m
           ON m.id = v.mp_id ORDER BY v.date_mouvement, v.id""").fetchall()
    _feuille(wb, "Mouvements matières",
             [("Date", 12), ("Matière", 32), ("Type", 12), ("Quantité", 10), ("Unité", 8),
              ("Note", 34)],
             [(v["date_mouvement"], v["nom"], v["type"], round(v["quantite"], 3), v["unite"],
               v["notes"] or "") for v in mvts])


def feuille_achats(wb, conn, debut, fin):
    rows = conn.execute(
        """SELECT a.*, (SELECT group_concat(designation || COALESCE(' (' || quantite || ' ' ||
                   COALESCE(unite, '') || ')', ''), ' ; ') FROM achat_lignes l
                   WHERE l.achat_id = a.id) AS detail
           FROM achats a WHERE a.date_achat BETWEEN ? AND ? ORDER BY a.date_achat, a.id""",
        (debut, fin)).fetchall()
    _feuille(wb, "Registre des achats",
             [("Date", 12), ("Fournisseur", 28), ("N° de facture", 16), ("Nature de l'achat", 56),
              ("Mode de règlement", 16), ("Montant", 13), ("Justificatif", 11)],
             [(a["date_achat"], a["fournisseur"], a["reference"] or "", a["detail"] or "",
               MODES_PAIEMENT.get(a["mode_paiement"], ""), a["montant_total"],
               "oui" if a["justificatif"] else "non") for a in rows], {6: EURO})


def feuille_clients(wb, conn):
    from vues_clients import STATS_CLIENTS
    rows = conn.execute(STATS_CLIENTS + " ORDER BY c.nom").fetchall()
    _feuille(wb, "Clients",
             [("Nom", 28), ("Type", 13), ("Email", 30), ("Téléphone", 15), ("Adresse", 36),
              ("Accepte les nouveautés", 12), ("Bouteilles achetées", 12), ("Total dépensé", 14),
              ("Dernier achat", 13), ("Bouteilles consignées", 12), ("Notes", 36)],
             [(c["nom"], TYPES_CLIENTS.get(c["type"], c["type"]), c["email"], c["telephone"],
               c["adresse"], "oui" if c["accepte_infos"] else "non", c["bouteilles"],
               c["depense"], c["dernier_achat"], c["consignes"], c["notes"]) for c in rows],
             {8: EURO})
    bons = conn.execute("SELECT * FROM bons ORDER BY date_emission").fetchall()
    _feuille(wb, "Bons d'achat",
             [("Code", 16), ("Bénéficiaire", 26), ("Montant", 11), ("Motif", 36),
              ("Émission", 12), ("Expiration", 12), ("Utilisé", 9)],
             [(b["code"], b["beneficiaire"], b["montant"], b["motif"], b["date_emission"],
               b["date_expiration"], "oui" if b["ticket_id"] else "non") for b in bons],
             {3: EURO})


def feuille_recettes(wb, conn):
    rows = conn.execute(
        """SELECT r.nom AS recette, r.duree_maceration_jours, r.degre, i.*
           FROM recettes r LEFT JOIN recette_ingredients i ON i.recette_id = r.id
           ORDER BY r.nom, i.id""").fetchall()
    _feuille(wb, "Recettes",
             [("Recette", 32), ("Macération (j)", 13), ("Degré", 8), ("Catégorie", 11),
              ("Ingrédient", 32), ("Quantité pour 1 L", 15), ("Unité", 8),
              ("Prix par unité", 13)],
             [(r["recette"], r["duree_maceration_jours"], r["degre"], r["categorie"], r["nom"],
               r["quantite_par_litre"], r["unite"], r["cout_unitaire"]) for r in rows])


EXPORTS = {
    "livre-recettes": "Livre des recettes",
    "achats": "Registre des achats",
    "ventes": "Ventes détaillées",
    "lots": "Registre des lots",
    "stock": "Stock",
    "clients": "Clients et bons d'achat",
    "recettes": "Recettes",
    "complet": "Classeur complet",
}
app.jinja_env.globals["EXPORTS"] = EXPORTS


@app.route("/exports/<nom>.xlsx")
def export(nom):
    if nom not in EXPORTS:
        abort(404)
    conn = get_db()
    conf = settings()
    couts = couts_par_lot(conn, conf)
    debut = parse_date(request.args.get("debut")) or "2000-01-01"
    fin = parse_date(request.args.get("fin")) or "2999-12-31"
    wb = Workbook()
    wb.remove(wb.active)
    if nom in ("livre-recettes", "complet"):
        feuille_livre(wb, conn, debut, fin)
    if nom in ("ventes", "complet"):
        feuille_ventes(wb, conn, debut, fin, couts)
    if nom in ("lots", "complet"):
        feuille_lots(wb, conn, conf)
    if nom in ("achats", "complet"):
        feuille_achats(wb, conn, debut, fin)
    if nom in ("stock", "complet"):
        feuille_stock(wb, conn, conf, couts)
        feuille_matieres(wb, conn)
    if nom in ("clients", "complet"):
        feuille_clients(wb, conn)
    if nom in ("recettes", "complet"):
        feuille_recettes(wb, conn)
    buf = BytesIO()
    wb.save(buf)
    fichier = f"MaisonCorsaire-{nom}-{datetime.now():%Y%m%d}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fichier}"'},
    )
